"""
AirHX DOCX report generator (python-docx).
"""
from __future__ import annotations
import io
from datetime import date

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from engines import (
    BundleDesignResult, BundleRatingResult, FluidProps,
    FIN_TYPE_LABELS,
)

# ── Colour palette ────────────────────────────────────────────────────────────
_DARK   = "1E3A5F"
_MID    = "2D5F8A"
_LIGHT  = "DCE4EF"
_ALT    = "F4F6FA"
_OK     = "D1FAE5"
_WARN   = "FEF3C7"
_FAIL   = "FEE2E2"
_BORDER = "B0B8C4"


# ── Low-level helpers ─────────────────────────────────────────────────────────

def _rgb(hex6: str) -> RGBColor:
    r, g, b = int(hex6[0:2], 16), int(hex6[2:4], 16), int(hex6[4:6], 16)
    return RGBColor(r, g, b)


def _set_cell_bg(cell, hex6: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd   = OxmlElement("w:shd")
    shd.set(qn("w:fill"),  hex6)
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:val"),   "clear")
    tc_pr.append(shd)


def _set_table_borders(table, color: str = _BORDER, sz: int = 4):
    tbl   = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"),   "single")
        el.set(qn("w:sz"),    str(sz))
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def _set_col_widths(table, widths_cm: list[float]):
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths_cm):
                cell.width = Cm(widths_cm[i])


def _heading(doc: Document, text: str, level: int = 2):
    p = doc.add_heading(text, level=level)
    p.runs[0].font.color.rgb = _rgb(_DARK if level == 1 else _MID)


def _sec_heading(doc: Document, letter: str, title: str):
    _heading(doc, f"{letter} — {title}", 2)


def _kv_table(doc: Document, pairs: list[tuple[str, str]]):
    table = doc.add_table(rows=len(pairs), cols=2)
    _set_table_borders(table)
    for i, (k, v) in enumerate(pairs):
        row = table.rows[i]
        row.cells[0].text = k
        row.cells[0].paragraphs[0].runs[0].bold = True
        row.cells[0].paragraphs[0].runs[0].font.color.rgb = _rgb(_MID)
        row.cells[1].text = v
        if i % 2 == 1:
            _set_cell_bg(row.cells[0], _ALT)
            _set_cell_bg(row.cells[1], _ALT)
    _set_col_widths(table, [6.5, 9.5])
    doc.add_paragraph()


# ── Main generator ────────────────────────────────────────────────────────────

def generate_word_report(
    proj_name: str, tag: str, issued_for: str,
    mode: str,
    result,
    fluid: FluidProps,
    hybrid_result=None,
    pump_result=None,
    exp_result=None,
    pipe_result=None,
) -> bytes:

    doc   = Document()
    today = date.today().strftime("%d %b %Y")
    g     = result.geom
    htc   = result.htc
    is_design = isinstance(result, BundleDesignResult)

    # Page margins
    for section in doc.sections:
        section.top_margin    = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin   = Cm(2.0)
        section.right_margin  = Cm(2.0)

    # ── Cover ────────────────────────────────────────────────────────────────
    doc.add_heading("AirHX — Air-Cooled Heat Exchanger", 1)
    doc.paragraphs[-1].runs[0].font.color.rgb = _rgb(_DARK)

    _kv_table(doc, [
        ("Project",       proj_name),
        ("Tag / unit",    tag),
        ("Issued for",    issued_for),
        ("Mode",          mode),
        ("Date",          today),
    ])

    doc.add_page_break()

    # ── A: Design basis ───────────────────────────────────────────────────────
    _sec_heading(doc, "A", "Design basis")
    _kv_table(doc, [
        ("Tube-side fluid",     fluid.name),
        ("Process inlet T",     f"{result.T_proc_in:.1f} °C"),
        ("Process outlet T",    f"{result.T_proc_out:.1f} °C"),
        ("Air inlet T",         f"{result.T_air_in:.1f} °C"),
        ("Air outlet T",        f"{result.T_air_out:.1f} °C" if hasattr(result, "T_air_out") else "—"),
        ("Heat duty Q",         f"{result.Q_kW:.0f} kW"),
        ("Process mass flow",   f"{result.m_proc_kgs:.2f} kg/s"),
    ])

    # ── B: Bundle geometry ────────────────────────────────────────────────────
    _sec_heading(doc, "B", "Bundle geometry")
    _kv_table(doc, [
        ("Fin type",            FIN_TYPE_LABELS.get(g.fin_type, g.fin_type)),
        ("Number of bays",      str(g.n_bays)),
        ("Rows per bay",        str(g.n_rows)),
        ("Tubes/row (per bay)", str(g.n_tubes_row)),
        ("Tube-side passes",    str(g.n_passes)),
        ("Tube length",         f"{g.L_tube_m:.3f} m"),
        ("Total tubes",         str(g.n_tubes_total)),
        ("Fan arrangement",     g.fan_type),
    ])

    # ── C: Thermal performance ────────────────────────────────────────────────
    _sec_heading(doc, "C", "Thermal performance")
    rows: list[tuple[str, str]] = [
        ("Overall U (total ext. area)",  f"{htc.U:.1f} W/m²K"),
        ("Air-side HTC (total area)",    f"{htc.h_air_eff:.1f} W/m²K"),
        ("Fin efficiency η_fin",         f"{htc.eta_fin*100:.1f}%"),
        ("Tube-side HTC",               f"{htc.h_tube:.0f} W/m²K"),
        ("LMTD (counter-flow)",         f"{htc.LMTD:.1f} K"),
        ("F-factor",                    f"{htc.F:.3f}"),
        ("Effective LMTD",              f"{htc.LMTD_eff:.1f} K"),
    ]
    if is_design:
        rows += [
            ("Required area",   f"{result.A_req_m2:.0f} m²"),
            ("Provided area",   f"{result.A_total_m2:.0f} m²"),
            ("Area margin",     f"{result.area_margin*100:.0f}%"),
        ]
    else:
        rows += [
            ("Total area",      f"{result.A_total_m2:.0f} m²"),
            ("Effectiveness ε", f"{result.epsilon:.3f}"),
        ]
    _kv_table(doc, rows)

    # ── D: Air side ───────────────────────────────────────────────────────────
    _sec_heading(doc, "D", "Air side")
    _kv_table(doc, [
        ("Air mass flow",       f"{result.m_air_kgs:.1f} kg/s"),
        ("Air mass flux G",     f"{result.G_air:.2f} kg/(m²·s)"),
        ("Air-side ΔP",         f"{result.dP_air_Pa:.0f} Pa"),
        ("Fan power",           f"{result.P_fan_kW:.1f} kW"),
        ("Tube velocity",       f"{result.v_tube_ms:.2f} m/s"),
        ("Tube-side Re",        f"{result.Re_tube:.0f}"),
    ])

    # ── E: Hybrid cooling ─────────────────────────────────────────────────────
    if hybrid_result is not None:
        _sec_heading(doc, "E", "Hybrid evaporative cooling")
        _kv_table(doc, [
            ("Mode",                    hybrid_result.mode.value),
            ("Wet-bulb temperature",    f"{hybrid_result.T_wb:.1f} °C"),
            ("Effective air inlet T",   f"{hybrid_result.T_air_in_eff:.1f} °C"),
            ("Dry duty",                f"{hybrid_result.Q_duty_dry_kW:.0f} kW"),
            ("Wet duty",                f"{hybrid_result.Q_duty_wet_kW:.0f} kW"),
            ("Water consumption",       f"{hybrid_result.m_water_m3h:.2f} m³/h"),
            ("Process outlet T (wet)",  f"{hybrid_result.T_proc_out_wet:.1f} °C"),
        ])

    # ── F: Cooling loop ───────────────────────────────────────────────────────
    if pump_result is not None or pipe_result is not None or exp_result is not None:
        _sec_heading(doc, "F", "Cooling loop sizing")
        loop_rows: list[tuple[str, str]] = []
        if pump_result:
            loop_rows += [
                ("Pump flow",               f"{pump_result.Q_m3h:.1f} m³/h"),
                ("Pump head",               f"{pump_result.H_m:.1f} m"),
                ("Shaft power",             f"{pump_result.P_shaft_kW:.1f} kW"),
                ("Motor size (std.)",        f"{pump_result.P_motor_kW:.0f} kW"),
            ]
        if pipe_result:
            loop_rows += [
                ("Main pipe",               f"DN{pipe_result.DN} — {pipe_result.OD_mm:.1f}×{pipe_result.t_wall_mm:.1f} mm"),
                ("Pipe velocity",           f"{pipe_result.v_ms:.2f} m/s"),
            ]
        if exp_result:
            loop_rows += [
                ("Expansion vessel",        f"{exp_result.V_vessel_L:.0f} L"),
                ("Pre-charge pressure",     f"{exp_result.p_charge_bara:.2f} bara"),
            ]
        _kv_table(doc, loop_rows)

    # ── G: Warnings ───────────────────────────────────────────────────────────
    _sec_heading(doc, "G", "Notes & findings")
    warn_items = (result.warnings or []) + (hybrid_result.notes if hybrid_result else [])
    if warn_items:
        for w in warn_items:
            p = doc.add_paragraph(f"⚠  {w}", style="List Bullet")
    else:
        doc.add_paragraph("No engineering warnings.")

    doc.add_paragraph(
        "AirHX is a screening tool only. Results are approximate and require "
        "verification by a qualified engineer. Not for certified design use."
    ).italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
